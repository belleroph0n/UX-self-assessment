"""One-off export of the UX Health Check criteria workbook (v1.51) into the
content files used by the app. After this export, the YAML files in /content
are the source of truth; the workbook is kept only as a historical reference.

Usage: python tools/export_workbook.py path/to/workbook.xlsx content/
"""
import sys, os, re
import yaml
from openpyxl import load_workbook

class Dumper(yaml.SafeDumper):
    pass
def str_presenter(dumper, data):
    if "\n" in data:
        return dumper.represent_scalar("tag:yaml.org,2002:str", data, style="|")
    return dumper.represent_scalar("tag:yaml.org,2002:str", data)
Dumper.add_representer(str, str_presenter)

def dump(path, header, data):
    with open(path, "w", encoding="utf-8") as f:
        f.write(header.rstrip() + "\n\n")
        yaml.dump(data, f, Dumper=Dumper, sort_keys=False, allow_unicode=True, width=1000)

def clean(v):
    if v is None:
        return ""
    if isinstance(v, str):
        return v.strip()
    return v

APPLIES = {
    "All": "all", "Public-facing": "public-facing", "Transactional": "transactional",
    "Authenticated": "authenticated", "Uses media": "uses-media", "Uses CMS": "uses-cms", "Uses AI": "uses-ai",
}
STATUS = {"Mandatory (NZ)": "mandatory", "Recommended (NZ)": "recommended", "Good practice": "good-practice"}

def main(xlsx, out):
    wb = load_workbook(xlsx, data_only=True)
    os.makedirs(out, exist_ok=True)

    # categories (from detailed Scoring tab)
    sc = wb["Scoring"]
    code_by_name = {}
    cats = []
    crit = wb["Criteria"]
    for r in range(2, crit.max_row + 1):
        cid = crit.cell(r, 1).value
        if cid:
            code_by_name[crit.cell(r, 2).value] = cid.split("-")[0]
    for r in range(18, 28):
        name = sc.cell(r, 2).value
        cats.append({"id": code_by_name[name], "name": name, "points": sc.cell(r, 3).value,
                     "description": clean(sc.cell(r, 9).value)})
    assert sum(c["points"] for c in cats) == 100
    dump(f"{out}/categories.yaml",
         "# Assessment categories. 'points' is each category's share of the 100-point UX health score\n"
         "# and must total 100. 'description' is shown to assessors as 'what good looks like'.",
         {"categories": cats})

    # criteria
    rows = []
    for r in range(2, crit.max_row + 1):
        cid = crit.cell(r, 1).value
        if not cid:
            continue
        v = lambda c: clean(crit.cell(r, c).value)
        item = {
            "id": cid,
            "category": cid.split("-")[0],
            "subcategory": v(3),
            "question": v(4),
            "evidence": v(6),
            "weight": int(v(7)),
            "critical": v(8) == "Y",
            "recommendation": v(9),
            "source": v(10),
            "source_version": v(11),
            "nz_status": STATUS[v(12)],
            "applies_to": APPLIES[v(13)],
            "link": {"label": v(23), "url": v(24)},
        }
        if v(14):
            item["review_note"] = v(14)
        rows.append(item)
    dump(f"{out}/criteria.yaml",
         "# Master list of criteria. Every criterion is one question in the detailed assessment.\n"
         "# Criteria with weight 3 also form the simple assessment (the app's default).\n"
         "# Field guide: see docs/editing-content.md",
         {"criteria": rows})

    # scoring parameters (simple set parameters - identical defaults to detailed)
    ss = wb["Scoring (simple)"]
    scoring = {
        "answer_points": {"yes": ss["C6"].value, "partial": ss["C7"].value, "no": ss["C8"].value},
        "critical_cap": ss["C10"].value,
        "ratings": [
            {"id": "healthy", "label": "Healthy", "min": ss["C11"].value, "colour": "green"},
            {"id": "good", "label": "Good, with gaps", "min": ss["C12"].value, "colour": "yellow"},
            {"id": "needs-attention", "label": "Needs attention", "min": ss["C13"].value, "colour": "amber"},
            {"id": "at-risk", "label": "At risk", "min": 0, "colour": "red"},
        ],
        "critical_failure_colour": "red",
        "key_recommendations": 5,
        "simple_set_weight": 3,
    }
    dump(f"{out}/scoring.yaml",
         "# Scoring parameters for the UX health score (both simple and detailed assessments).\n"
         "# ratings: a score at or above 'min' gets that rating; list from highest to lowest.\n"
         "# Any critical criterion answered 'no' caps the score at critical_cap and shows red.",
         scoring)

    # materiality
    mt = wb["Materiality"]
    factors = []
    for r in range(5, 15):
        factors.append({
            "id": mt.cell(r, 1).value, "factor": clean(mt.cell(r, 2).value), "question": clean(mt.cell(r, 3).value),
            "anchors": [clean(mt.cell(r, c).value) for c in range(4, 9)],
            "weight": mt.cell(r, 9).value, "floor_trigger": mt.cell(r, 10).value == "Y",
        })
    levels = []
    colours = ["green", "yellow-green", "yellow", "orange", "red"]
    for k in range(5):
        r = 31 + k
        levels.append({"rating": mt.cell(r, 2).value, "label": clean(mt.cell(r, 3).value),
                       "meaning": clean(mt.cell(r, 4).value), "colour": colours[k]})
    dump(f"{out}/materiality.yaml",
         "# Materiality: how much the application matters, rated 1 (Immaterial) to 5 (Critical).\n"
         "# Scored separately from the UX health score and never combined with it.\n"
         "# Rating = weighted average of answered factors, rounded; if any floor_trigger factor\n"
         "# is answered 5, the rating can't be lower than floor_minimum.",
         {"floor_minimum": mt["D18"].value, "levels": levels, "factors": factors})

    # tiers
    tt = wb["Service tiers"]
    keys = ["id", "materiality", "when_to_use", "availability", "support_hours", "p1_incident", "p2_incident",
            "rto", "rpo", "monitoring", "maintenance", "health_check_frequency", "target_rating_text",
            "remediation", "reporting"]
    tiers = []
    for r in range(5, 9):
        t = {k: clean(tt.cell(r, i + 1).value) for i, k in enumerate(keys)}
        t["name"] = t["id"]; t["id"] = t["id"].lower()
        tiers.append(t)
    target = {"platinum": "healthy", "gold": "good", "silver": "needs-attention", "bronze": "at-risk"}
    for t in tiers:
        t["target_rating"] = target[t["id"]]
    rating_to_tier = {tt.cell(r, 1).value: tt.cell(r, 2).value.lower() for r in range(12, 17)}
    labels = ["Tier", "Materiality rating", "When to use", "Availability target", "Support hours",
              "Priority 1 incident: response / restore", "Priority 2 incident: response / restore",
              "Recovery time objective (RTO)", "Recovery point objective (RPO)", "Monitoring and alerting",
              "Planned maintenance", "UX health check frequency", "Target UX health rating",
              "Remediation of health check findings", "Performance reporting"]
    dump(f"{out}/tiers.yaml",
         "# Proposed service tier definitions (not yet agreed by the organisation).\n"
         "# rating_to_tier maps the materiality rating (1-5) to a suggested tier.\n"
         "# target_rating is the minimum UX health rating expected for the tier.",
         {"status": "proposed", "rating_to_tier": rating_to_tier,
          "field_labels": dict(zip(["name", "materiality", "when_to_use", "availability", "support_hours",
                                    "p1_incident", "p2_incident", "rto", "rpo", "monitoring", "maintenance",
                                    "health_check_frequency", "target_rating_text", "remediation", "reporting"], labels)),
          "tiers": tiers})

    # sources
    so = wb["Sources"]
    srcs = []
    for r in range(2, so.max_row + 1):
        name, pub = so.cell(r, 1).value, so.cell(r, 2).value
        if not name or not pub:
            continue
        srcs.append({"name": clean(name), "publisher": clean(pub), "status": clean(so.cell(r, 3).value),
                     "nz_status": clean(so.cell(r, 4).value), "url": clean(so.cell(r, 5).value),
                     "link_check": clean(so.cell(r, 6).value), "licence": clean(so.cell(r, 7).value),
                     "used_in": clean(so.cell(r, 8).value)})
    dump(f"{out}/sources.yaml",
         "# Register of standards, laws, guidance and tools referenced by the criteria.\n"
         "# Links are checked automatically by the link-check workflow.",
         {"sources": srcs})

    # version history -> CHANGELOG
    vs = wb["Version"]
    heads = [c.value for c in vs[1]]
    lines = ["# Content changelog", "",
             "History of the criteria library, carried over from the Version tab of the original workbook.",
             "Each release records what changed and why.", ""]
    latest = None
    for r in range(vs.max_row, 1, -1):
        row = [vs.cell(r, c).value for c in range(1, len(heads) + 1)]
        if not row[0]:
            continue
        latest = latest or row[0]
        lines.append(f"## {row[0]} - {row[1]}")
        lines.append("")
        for h, v in zip(heads[2:], row[2:]):
            if v:
                lines.append(f"### {h}")
                lines.append("")
                lines.append(str(v).strip())
                lines.append("")
    with open(os.path.join(os.path.dirname(out.rstrip("/")), "CHANGELOG.md"), "w", encoding="utf-8") as f:
        f.write("\n".join(lines))

    dump(f"{out}/meta.yaml",
         "# Content release details. Bump content_version for every change to the content files\n"
         "# and add an entry to CHANGELOG.md.",
         {"title": "UX Health Check", "content_version": latest.lstrip("v"), "released": "2026-10-24",
          "audience": "NZ Government agencies", "licence": "CC BY 4.0"})
    print("criteria", len(rows), "sources", len(srcs), "latest", latest)

if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
